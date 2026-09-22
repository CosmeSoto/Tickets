import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 Widescreen dimensions (13.333" x 7.5")
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Colors
    BG_DARK = RGBColor(15, 17, 23)        # Sleek slate/black #0F1117
    CARD_DARK = RGBColor(26, 32, 44)      # Dark card #1A202C
    CARD_BORDER = RGBColor(45, 55, 72)    # Border #2D3748
    TEXT_WHITE = RGBColor(255, 255, 255)  # Crisp white
    TEXT_MUTED = RGBColor(160, 174, 192)  # Muted silver/gray #A0AEC0
    TEXT_GOLD = RGBColor(212, 175, 55)    # Gold accent #D4AF37
    CYAN_ACCENT = RGBColor(56, 189, 248)  # Cyan #38BDF8
    GREEN_ACCENT = RGBColor(52, 211, 153) # Emerald #34D399
    PURPLE_ACCENT = RGBColor(192, 132, 252) # Violet #C084FC
    ORANGE_ACCENT = RGBColor(251, 146, 60)  # Warm orange #FB923C

    logo_path = "/Users/cosmesoto/.gemini/antigravity-ide/brain/03659b91-5e95-4161-bbcb-d0756881da02/.user_uploaded/media_1789999192289.png"
    banner_path = "/Users/cosmesoto/Documents/Tickets/sistema-tickets-nextjs/docs/marketing/psf_email_banner_1789999404275.jpg"
    poster_path = "/Users/cosmesoto/Documents/Tickets/sistema-tickets-nextjs/docs/marketing/psf_poster_adopcion_1789999436606.jpg"
    paco_banner_path = "/Users/cosmesoto/Documents/Tickets/sistema-tickets-nextjs/docs/marketing/paco_telegram_banner.jpg"
    paco_bot_path = "/Users/cosmesoto/Documents/Tickets/sistema-tickets-nextjs/docs/marketing/paco_bot.jpg"

    def set_slide_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_header(slide, category, title, subtitle=None):
        # Category Tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(10), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.name = 'Georgia'
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = TEXT_GOLD

        # Slide Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(10), Inches(0.65))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title
        p_title.font.name = 'Georgia'
        p_title.font.size = Pt(25)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE

        if subtitle:
            sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.45), Inches(11.5), Inches(0.4))
            tf_sub = sub_box.text_frame
            tf_sub.word_wrap = True
            tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = subtitle
            p_sub.font.name = 'Calibri'
            p_sub.font.size = Pt(13)
            p_sub.font.color.rgb = TEXT_MUTED

        # Add logo on top right if exists
        if os.path.exists(logo_path):
            slide.shapes.add_picture(logo_path, Inches(10.8), Inches(0.45), width=Inches(1.8))

    def add_card(slide, left, top, width, height, bg_color=CARD_DARK, border_color=CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        return shape

    # ==========================================
    # SLIDE 1: PORTADA
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, BG_DARK)

    if os.path.exists(banner_path):
        s1.shapes.add_picture(banner_path, Inches(6.2), Inches(1.0), width=Inches(6.5))

    tb_title = s1.shapes.add_textbox(Inches(0.9), Inches(1.4), Inches(5.2), Inches(4.8))
    tf1 = tb_title.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "PASEO SAN FRANCISCO"
    p0.font.name = 'Georgia'
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = TEXT_GOLD
    p0.space_after = Pt(12)

    p1 = tf1.add_paragraph()
    p1.text = "Transformación Digital &\nEficiencia Operativa"
    p1.font.name = 'Georgia'
    p1.font.size = Pt(30)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.space_after = Pt(14)

    p2 = tf1.add_paragraph()
    p2.text = "Lanzamiento de la Plataforma Institucional\nTickets · Noticias · Documentos · PACO Bot"
    p2.font.name = 'Calibri'
    p2.font.size = Pt(15)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_after = Pt(24)

    p3 = tf1.add_paragraph()
    p3.text = "Presentación Estratégica para Dirección y Jefaturas"
    p3.font.name = 'Calibri'
    p3.font.size = Pt(12)
    p3.font.bold = True
    p3.font.color.rgb = CYAN_ACCENT

    # ==========================================
    # SLIDE 2: EL DIAGNÓSTICO (PROBLEMA ACTUAL)
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, BG_DARK)
    add_header(s2, "Contexto y Oportunidad", "El Reto de la Gestión Descentralizada", "¿Por qué transformar nuestros canales tradicionales de atención y comunicación?")

    problems = [
        ("SOLICITUDES DISPERSAS", "Falta de Trazabilidad", "Requerimientos perdidos en WhatsApp, pasillos o correos personales. Sin responsables definidos ni tiempos de respuesta garantizados.", RGBColor(239, 68, 68)),
        ("FORMATOS OBSOLETOS", "Riesgo Administrativo", "Colaboradores utilizando versiones antiguas de formatos y políticas. Retrabajos innecesarios y falta de un repositorio oficial único.", RGBColor(245, 158, 11)),
        ("COMUNICACIÓN SATURADA", "Bajo Alcance Institucional", "Avisos corporativos que se pierden entre decenas de correos diarios. Imposible medir quién leyó las normativas y novedades.", RGBColor(168, 85, 247))
    ]

    card_w = Inches(3.64)
    card_h = Inches(4.3)
    card_top = Inches(2.0)

    for i, (tag, title, desc, tag_col) in enumerate(problems):
        c_left = Inches(0.8) + i * (card_w + Inches(0.4))
        add_card(s2, c_left, card_top, card_w, card_h)

        tb = s2.shapes.add_textbox(c_left + Inches(0.3), card_top + Inches(0.3), card_w - Inches(0.6), card_h - Inches(0.6))
        tf = tb.text_frame
        tf.word_wrap = True

        p_tag = tf.paragraphs[0]
        p_tag.text = tag
        p_tag.font.name = 'Calibri'
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = tag_col
        p_tag.space_after = Pt(10)

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.name = 'Georgia'
        p_t.font.size = Pt(18)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE
        p_t.space_after = Pt(14)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.name = 'Calibri'
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = TEXT_MUTED

    # Bottom Callout bar
    add_card(s2, Inches(0.8), Inches(6.45), Inches(11.73), Inches(0.6), bg_color=RGBColor(30, 41, 59), border_color=None)
    tb_bar = s2.shapes.add_textbox(Inches(1.0), Inches(6.5), Inches(11.3), Inches(0.5))
    p_b = tb_bar.text_frame.paragraphs[0]
    p_b.text = "🎯 Meta Principal: Centralizar el 100% de las solicitudes, documentos y avisos en un ecosistema auditable y en tiempo real."
    p_b.font.name = 'Calibri'
    p_b.font.size = Pt(13)
    p_b.font.bold = True
    p_b.font.color.rgb = TEXT_GOLD

    # ==========================================
    # SLIDE 3: LA SOLUCIÓN (LOS PILARES DEL SISTEMA)
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, BG_DARK)
    add_header(s3, "La Solución Integral", "Nuestra Plataforma Digital Institucional", "Un ecosistema unificado para optimizar el servicio, la información, los documentos y las alertas")

    pillars = [
        ("MÓDULO 1", "Mesa de Ayuda & Tickets", "• Registro ágil de solicitudes e incidencias\n• Prioridades (Urgente, Alta, Media, Baja)\n• Asignación automática por departamento\n• Trazabilidad y métricas SLA en vivo", CYAN_ACCENT),
        ("MÓDULO 2", "Noticias & Avisos", "• Muro institucional de comunicados\n• Segmentación por áreas y departamentos\n• Notificaciones en tiempo real\n• Confirmación y métricas de lectura", GREEN_ACCENT),
        ("MÓDULO 3", "Gestor Documental", "• Única fuente oficial de verdad (Formatos)\n• Descarga de plantillas oficiales vigentes\n• Búsqueda instantánea con filtros\n• Control de acceso y versiones oficiales", PURPLE_ACCENT)
    ]

    for i, (tag, title, desc, col) in enumerate(pillars):
        c_left = Inches(0.8) + i * (card_w + Inches(0.4))
        add_card(s3, c_left, card_top, card_w, card_h)

        tb = s3.shapes.add_textbox(c_left + Inches(0.3), card_top + Inches(0.3), card_w - Inches(0.6), card_h - Inches(0.6))
        tf = tb.text_frame
        tf.word_wrap = True

        p_tag = tf.paragraphs[0]
        p_tag.text = tag
        p_tag.font.name = 'Calibri'
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = col
        p_tag.space_after = Pt(10)

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.name = 'Georgia'
        p_t.font.size = Pt(18)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE
        p_t.space_after = Pt(14)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.name = 'Calibri'
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = TEXT_MUTED

    # Bottom bar highlighting multichannel notifications
    add_card(s3, Inches(0.8), Inches(6.45), Inches(11.73), Inches(0.6), bg_color=RGBColor(30, 41, 59), border_color=None)
    tb_bar = s3.shapes.add_textbox(Inches(1.0), Inches(6.5), Inches(11.3), Inches(0.5))
    p_b = tb_bar.text_frame.paragraphs[0]
    p_b.text = "🔔 Notificaciones Multicanal Integradas: Alertas en tiempo real vía Web (Campanita), Correo Electrónico y Bot de Telegram PACO."
    p_b.font.name = 'Calibri'
    p_b.font.size = Pt(12)
    p_b.font.bold = True
    p_b.font.color.rgb = CYAN_ACCENT

    # ==========================================
    # SLIDE 4: DETALLE TICKETS
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, BG_DARK)
    add_header(s4, "Profundización Módulo 1", "Mesa de Ayuda y Tickets de Soporte", "Control total sobre requerimientos, tiempos de respuesta y carga de trabajo")

    w_col = Inches(5.65)
    h_col = Inches(4.8)
    add_card(s4, Inches(0.8), Inches(2.0), w_col, h_col)
    tb_l = s4.shapes.add_textbox(Inches(1.1), Inches(2.3), w_col - Inches(0.6), h_col - Inches(0.6))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True

    p = tf_l.paragraphs[0]
    p.text = "CARACTERÍSTICAS OPERATIVAS"
    p.font.name = 'Calibri'
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT
    p.space_after = Pt(12)

    features = [
        ("Flujo de Estados Transparente", "Abierto → En Progreso → Resuelto → Cerrado."),
        ("Acuerdos de Nivel de Servicio (SLA)", "Tiempos máximos de respuesta según urgencia."),
        ("Asignación Inteligente", "Direccionamiento automático al técnico o área correspondiente."),
        ("Historial & Colaboración", "Registro auditable de comentarios, adjuntos y tiempos."),
        ("Base de Conocimientos", "Casos resueltos se convierten en guías para acelerar soporte.")
    ]
    for ft, fd in features:
        p_ft = tf_l.add_paragraph()
        p_ft.text = f"✔ {ft}: "
        p_ft.font.name = 'Calibri'
        p_ft.font.size = Pt(13)
        p_ft.font.bold = True
        p_ft.font.color.rgb = TEXT_WHITE
        run = p_ft.add_run()
        run.text = fd
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED
        p_ft.space_after = Pt(8)

    # Right Column: Value for Management
    add_card(s4, Inches(6.85), Inches(2.0), w_col, h_col)
    tb_r = s4.shapes.add_textbox(Inches(7.15), Inches(2.3), w_col - Inches(0.6), h_col - Inches(0.6))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True

    p = tf_r.paragraphs[0]
    p.text = "VALOR PARA JEFATURAS Y DIRECCIÓN"
    p.font.name = 'Calibri'
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = TEXT_GOLD
    p.space_after = Pt(12)

    mgmt_val = [
        ("Métricas de Rendimiento Reales", "Saber con precisión cuántas solicitudes atiende cada área al mes."),
        ("Detección de Cuellos de Botella", "Identificar qué procesos o equipos generan más incidencias."),
        ("Exportación Ejecutiva", "Descarga de reportes detallados en Excel y PDF para comités."),
        ("Eliminación de la Informalidad", "Fin a las solicitudes perdidas en pasillos o chats personales."),
        ("Mejora en la Satisfacción Interna", "El colaborador siempre sabe quién atiende su caso.")
    ]
    for ft, fd in mgmt_val:
        p_ft = tf_r.add_paragraph()
        p_ft.text = f"★ {ft}: "
        p_ft.font.name = 'Calibri'
        p_ft.font.size = Pt(13)
        p_ft.font.bold = True
        p_ft.font.color.rgb = TEXT_GOLD
        run = p_ft.add_run()
        run.text = fd
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED
        p_ft.space_after = Pt(8)

    # ==========================================
    # SLIDE 5: NOTICIAS Y DOCUMENTOS
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, BG_DARK)
    add_header(s5, "Profundización Módulos 2 y 3", "Noticias Oficiales y Gestor Documental", "Comunicación segmentada sin saturación y control estricto de formatos vigentes")

    # Left: Noticias
    add_card(s5, Inches(0.8), Inches(2.0), w_col, h_col)
    tb_n = s5.shapes.add_textbox(Inches(1.1), Inches(2.3), w_col - Inches(0.6), h_col - Inches(0.6))
    tf_n = tb_n.text_frame
    tf_n.word_wrap = True

    p = tf_n.paragraphs[0]
    p.text = "NOTICIAS & COMUNICACIÓN INTERNA"
    p.font.name = 'Calibri'
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = GREEN_ACCENT
    p.space_after = Pt(12)

    news_points = [
        ("Segmentación por Familias/Áreas", "Publicaciones generales o exclusivas para departamentos específicos."),
        ("Notificaciones en Tiempo Real", "Alertas inmediatas en plataforma, correo y Bot de Telegram."),
        ("Muro Multimedia Dinámico", "Soporte para imágenes, comunicados oficiales y documentos adjuntos."),
        ("Trazabilidad de Alcance", "Control exacto de vistas y confirmaciones de lectura institucional.")
    ]
    for ft, fd in news_points:
        p_ft = tf_n.add_paragraph()
        p_ft.text = f"📰 {ft}\n"
        p_ft.font.name = 'Calibri'
        p_ft.font.size = Pt(13)
        p_ft.font.bold = True
        p_ft.font.color.rgb = TEXT_WHITE
        run = p_ft.add_run()
        run.text = f"   {fd}"
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED
        p_ft.space_after = Pt(10)

    # Right: Documentos
    add_card(s5, Inches(6.85), Inches(2.0), w_col, h_col)
    tb_d = s5.shapes.add_textbox(Inches(7.15), Inches(2.3), w_col - Inches(0.6), h_col - Inches(0.6))
    tf_d = tb_d.text_frame
    tf_d.word_wrap = True

    p = tf_d.paragraphs[0]
    p.text = "GESTOR DOCUMENTAL Y FORMATOS"
    p.font.name = 'Calibri'
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = PURPLE_ACCENT
    p.space_after = Pt(12)

    doc_points = [
        ("Única Fuente Oficial (Single Source)", "Los colaboradores descargan siempre la versión vigente del formato."),
        ("Categorización Estructurada", "Formatos de RRHH, Compras, Operaciones, TI y Mantenimiento."),
        ("Búsqueda Instantánea", "Localización de plantillas y manuales en menos de 5 segundos."),
        ("Seguridad y Permisos", "Visibilidad restringida según el rol y departamento del usuario.")
    ]
    for ft, fd in doc_points:
        p_ft = tf_d.add_paragraph()
        p_ft.text = f"📂 {ft}\n"
        p_ft.font.name = 'Calibri'
        p_ft.font.size = Pt(13)
        p_ft.font.bold = True
        p_ft.font.color.rgb = TEXT_WHITE
        run = p_ft.add_run()
        run.text = f"   {fd}"
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED
        p_ft.space_after = Pt(10)

    # ==========================================
    # SLIDE 6: NUEVA SLIDE - NOTIFICACIONES MULTICANAL & PACO BOT
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, BG_DARK)
    add_header(s6, "Innovación y Conectividad", "Ecosistema de Notificaciones Multicanal", "Garantizando que ninguna solicitud ni aviso quede sin atender en tiempo real")

    # Left: 3 Channels cards
    w_notif = Inches(7.5)
    channels = [
        ("🔔 CAMPANITA WEB (IN-APP)", "Alertas en Tiempo Real en la Plataforma", "Actualizaciones instantáneas mientras el usuario navega (Server-Sent Events). Contador de notificaciones no leídas y acceso rápido al ticket o aviso.", CYAN_ACCENT),
        ("📧 CORREO ELECTRÓNICO", "Notificaciones Formales y Resúmenes", "Envío automático de confirmaciones de apertura, asignación de responsable, cambios de estado y comunicados corporativos directamente a la bandeja.", GREEN_ACCENT),
        ("🤖 PACO — BOT DE TELEGRAM", "Asistente Móvil de Paseo San Francisco", "Notificaciones push instantáneas al celular del colaborador o técnico. Ideal para personal operativo y de campo que no está permanentemente frente a una PC.", ORANGE_ACCENT)
    ]

    for i, (tag, title, desc, accent_col) in enumerate(channels):
        c_top = Inches(2.0) + i * Inches(1.6)
        add_card(s6, Inches(0.8), c_top, w_notif, Inches(1.45))

        tb = s6.shapes.add_textbox(Inches(1.1), c_top + Inches(0.18), w_notif - Inches(0.6), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = f"{tag}  —  {title}"
        p_t.font.name = 'Georgia'
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = accent_col
        p_t.space_after = Pt(4)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.name = 'Calibri'
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = TEXT_MUTED

    # Right: PACO Bot Feature Card with Image
    w_paco = Inches(4.0)
    h_paco = Inches(4.8)
    add_card(s6, Inches(8.55), Inches(2.0), w_paco, h_paco)

    if os.path.exists(paco_bot_path):
        s6.shapes.add_picture(paco_bot_path, Inches(8.8), Inches(2.2), width=Inches(3.5))

    tb_paco = s6.shapes.add_textbox(Inches(8.75), Inches(4.85), Inches(3.6), Inches(1.8))
    tf_p = tb_paco.text_frame
    tf_p.word_wrap = True

    p_p0 = tf_p.paragraphs[0]
    p_p0.text = "PACO: Bot Oficial PSF"
    p_p0.font.name = 'Georgia'
    p_p0.font.size = Pt(15)
    p_p0.font.bold = True
    p_p0.font.color.rgb = TEXT_GOLD
    p_p0.alignment = PP_ALIGN.CENTER
    p_p0.space_after = Pt(4)

    p_p1 = tf_p.add_paragraph()
    p_p1.text = "Conecta a técnicos y colaboradores con alertas críticas en segundos desde Telegram."
    p_p1.font.name = 'Calibri'
    p_p1.font.size = Pt(12)
    p_p1.font.color.rgb = TEXT_WHITE
    p_p1.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 7: KPIS Y RESULTADOS ESPERADOS
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7, BG_DARK)
    add_header(s7, "Impacto y Métricas de Éxito", "Resultados Proyectados a 90 Días", "Objetivos claros para medir el retorno de la inversión y la mejora en eficiencia")

    kpis = [
        ("-40%", "Tiempo de Respuesta", "Reducción drástica en el tiempo desde que surge una solicitud hasta la primera atención.", CYAN_ACCENT),
        ("100%", "Trazabilidad Total", "Cero requerimientos 'perdidos'; cada solicitud cuenta con folio, historial y responsable.", GREEN_ACCENT),
        ("< 5%", "Canales Informales", "Disminución de solicitudes por WhatsApp y pasillos hacia la plataforma oficial.", TEXT_GOLD),
        ("> 85%", "Tasa de Adopción", "Participación activa de los colaboradores en el primer mes de operación.", PURPLE_ACCENT)
    ]

    kw = Inches(2.68)
    kh = Inches(4.5)
    for i, (stat, title, desc, col) in enumerate(kpis):
        c_left = Inches(0.8) + i * (kw + Inches(0.33))
        add_card(s7, c_left, Inches(2.1), kw, kh)

        tb = s7.shapes.add_textbox(c_left + Inches(0.2), Inches(2.4), kw - Inches(0.4), kh - Inches(0.6))
        tf = tb.text_frame
        tf.word_wrap = True

        p_s = tf.paragraphs[0]
        p_s.text = stat
        p_s.font.name = 'Georgia'
        p_s.font.size = Pt(36)
        p_s.font.bold = True
        p_s.font.color.rgb = col
        p_s.alignment = PP_ALIGN.CENTER
        p_s.space_after = Pt(14)

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.name = 'Georgia'
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE
        p_t.alignment = PP_ALIGN.CENTER
        p_t.space_after = Pt(14)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.name = 'Calibri'
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = TEXT_MUTED
        p_d.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 8: PLAN DE DESPLIEGUE ACTUALIZADO POR EL USUARIO
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8, BG_DARK)
    add_header(s8, "Estrategia de Lanzamiento", "Plan de Adopción y Cronograma", "Asegurando una transición fluida y una alta participación de los colaboradores")

    phases = [
        ("FASE 1: PREPARACIÓN", "Validación & Ajustes\n(Días 1 - 5)", "• Validación de roles y permisos con jefes de área.\n• Carga final de formatos oficiales vigentes.\n• Configuración de categorías de tickets y SLAs."),
        ("FASE 2: SOCIALIZACIÓN", "Campaña de Adopción\n(Días 7 - 15)", "• Envío de correos de expectativa (Teasers).\n• Publicación de afiches con código QR en áreas clave.\n• Micro-capacitaciones prácticas de 15 minutos mediante medios de video en comunicaciones."),
        ("FASE 3: GO-LIVE", "Operación & Monitoreo\n(Día 30+)", "• Apertura oficial del portal para toda la organización.\n• Mesa de soporte activa para resolver dudas iniciales.\n• Entrega del primer reporte ejecutivo de métricas.")
    ]

    pw = Inches(3.64)
    ph = Inches(4.5)
    for i, (tag, title, desc) in enumerate(phases):
        c_left = Inches(0.8) + i * (pw + Inches(0.4))
        add_card(s8, c_left, Inches(2.1), pw, ph)

        tb = s8.shapes.add_textbox(c_left + Inches(0.3), Inches(2.3), pw - Inches(0.6), ph - Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p_tag = tf.paragraphs[0]
        p_tag.text = tag
        p_tag.font.name = 'Calibri'
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = TEXT_GOLD
        p_tag.space_after = Pt(8)

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.name = 'Georgia'
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE
        p_t.space_after = Pt(14)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.name = 'Calibri'
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 9: CIERRE Y CONCLUSIÓN
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9, BG_DARK)

    if os.path.exists(poster_path):
        s9.shapes.add_picture(poster_path, Inches(7.5), Inches(1.0), height=Inches(5.5))

    tb_c = s9.shapes.add_textbox(Inches(0.9), Inches(1.6), Inches(6.0), Inches(4.5))
    tfc = tb_c.text_frame
    tfc.word_wrap = True

    p0 = tfc.paragraphs[0]
    p0.text = "PASEO SAN FRANCISCO"
    p0.font.name = 'Georgia'
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = TEXT_GOLD
    p0.space_after = Pt(14)

    p1 = tfc.add_paragraph()
    p1.text = "Listos para el Lanzamiento Oficial"
    p1.font.name = 'Georgia'
    p1.font.size = Pt(30)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.space_after = Pt(18)

    p2 = tfc.add_paragraph()
    p2.text = "Una plataforma diseñada para optimizar los tiempos de respuesta, garantizar el orden operativo y conectar a toda la organización con alertas multicanal."
    p2.font.name = 'Calibri'
    p2.font.size = Pt(15)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_after = Pt(28)

    p3 = tfc.add_paragraph()
    p3.text = "¿Preguntas o Comentarios de la Dirección?"
    p3.font.name = 'Georgia'
    p3.font.size = Pt(18)
    p3.font.bold = True
    p3.font.color.rgb = CYAN_ACCENT

    # Save presentation
    output_dir = "/Users/cosmesoto/Documents/Tickets/sistema-tickets-nextjs/docs"
    os.makedirs(output_dir, exist_ok=True)
    out_path = os.path.join(output_dir, "Presentacion_Lanzamiento_Paseo_San_Francisco.pptx")
    prs.save(out_path)
    print(f"Presentation saved successfully at: {out_path}")

    # Also save in brain artifacts directory
    artifact_dir = "/Users/cosmesoto/.gemini/antigravity-ide/brain/03659b91-5e95-4161-bbcb-d0756881da02"
    artifact_out = os.path.join(artifact_dir, "Presentacion_Lanzamiento_Paseo_San_Francisco.pptx")
    prs.save(artifact_out)
    print(f"Artifact presentation saved at: {artifact_out}")

if __name__ == "__main__":
    create_presentation()
